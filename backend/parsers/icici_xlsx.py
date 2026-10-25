"""ICICI credit-card statement parser (Excel export).

ICICI's "Download as Excel" gives an OOXML .xlsx file — often saved with a .xls name, which openpyxl
refuses by extension, so we read the bytes through BytesIO. Layout (sheet 'CCStatement'):

    ... header rows (card details, Statement Period, totals) ...
    Transaction Date | Details | Amount (INR) | Reward Points | Reference Number
    13-06-2026       | SWIGGY  | 229 Dr.      | 2            | 40012345678

Amounts carry a 'Dr.'/'Cr.' suffix → debit/credit. The reference number is the natural dedupe key.
"""

from __future__ import annotations

import io
import re
from datetime import datetime

from parsers.base import ParsedTransaction, compute_stmt_month

_DATE_RE = re.compile(r"^\d{2}-\d{2}-\d{4}$")


def _load_sheet(path: str):
    data = open(path, "rb").read()
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    name = "CCStatement" if "CCStatement" in wb.sheetnames else wb.sheetnames[0]
    return wb[name]


def parse_icici_xlsx(path: str, card_type: str, card_last4: str) -> list[ParsedTransaction]:
    ws = _load_sheet(path)
    out: list[ParsedTransaction] = []
    started = False
    for row in ws.iter_rows(values_only=True):
        cs = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
        if not cs:
            continue
        if not started:
            if any("Transaction Date" in c for c in cs):
                started = True
            continue
        if not _DATE_RE.match(cs[0]):
            continue
        amt_field = next((c for c in cs if c.endswith("Dr.") or c.endswith("Cr.")), None)
        if not amt_field:
            continue
        try:
            amount = float(amt_field.rsplit(" ", 1)[0].replace(",", ""))
        except ValueError:
            continue
        direction = "debit" if amt_field.endswith("Dr.") else "credit"
        iso = datetime.strptime(cs[0], "%d-%m-%Y").strftime("%Y-%m-%d")
        out.append(ParsedTransaction(
            bank="ICICI", card_type=card_type or "Credit Card", card_last4=card_last4 or "",
            stmt_month=compute_stmt_month(iso), date=iso, description=cs[1] if len(cs) > 1 else "",
            amount=amount, direction=direction,
        ))
    return out
